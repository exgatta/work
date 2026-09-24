"""sources/ の資料を sources_text/ にテキスト化する（引用用のページ・スライド番号つき）。

使い方: ノートブックのフォルダで `python3 convert.py`
対応: .pdf .docx .pptx .xlsx .txt .md .csv
"""
import subprocess
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
SRC = HERE / "sources"
OUT = HERE / "sources_text"


def ensure_libs():
    try:
        import docx, openpyxl, pptx, pypdf  # noqa: F401
    except Exception:
        # クラウドのコンテナは毎回まっさらなので、足りなければ入れる
        subprocess.run(
            [sys.executable, "-m", "pip", "install", "-q",
             "cffi", "pypdf", "python-docx", "python-pptx", "openpyxl"],
            check=True,
        )


def pdf(path):
    from pypdf import PdfReader
    parts, empty = [], 0
    for i, page in enumerate(PdfReader(path).pages, 1):
        text = (page.extract_text() or "").strip()
        if not text:
            empty += 1
        parts.append(f"## p.{i}\n\n{text}")
    note = "（文字を取り出せないページあり。スキャン画像の可能性 → PDF を直接読んで確認する）\n\n" if empty else ""
    return note + "\n\n".join(parts)


def docx_(path):
    import docx
    d = docx.Document(path)
    lines = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n\n".join(lines)


def pptx_(path):
    from pptx import Presentation
    parts = []
    for i, slide in enumerate(Presentation(path).slides, 1):
        texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            texts.append("[ノート] " + slide.notes_slide.notes_text_frame.text)
        parts.append(f"## スライド{i}\n\n" + "\n".join(texts))
    return "\n\n".join(parts)


def xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    parts = []
    for ws in wb.worksheets:
        rows = [" | ".join("" if v is None else str(v) for v in r)
                for r in ws.iter_rows(values_only=True) if any(v is not None for v in r)]
        parts.append(f"## シート「{ws.title}」\n\n" + "\n".join(rows))
    return "\n\n".join(parts)


def plain(path):
    return path.read_text(encoding="utf-8", errors="replace")


CONVERTERS = {".pdf": pdf, ".docx": docx_, ".pptx": pptx_, ".xlsx": xlsx,
              ".txt": plain, ".md": plain, ".csv": plain}


def main():
    ensure_libs()
    OUT.mkdir(exist_ok=True)
    for path in sorted(SRC.rglob("*")):
        if not path.is_file() or path.name.startswith("."):
            continue
        conv = CONVERTERS.get(path.suffix.lower())
        rel = path.relative_to(SRC)
        if conv is None:
            print(f"対象外（画像などは直接読む）: {rel}")
            continue
        out = OUT / rel.with_suffix(rel.suffix + ".md")
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(f"# 出典: sources/{rel}\n\n{conv(path)}\n", encoding="utf-8")
        print(f"変換: {rel} -> sources_text/{out.relative_to(OUT)}")


if __name__ == "__main__":
    main()
